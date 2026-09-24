from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.mail import send_mail
from django.db import IntegrityError, transaction
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncDate
from django.db import connection
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings
import csv, hashlib, hmac, io, json, re, secrets, urllib.request
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from .forms import ApiKeyCreateForm, BulkDataImportForm, CommunicationForm, ConsultantCreateForm, ContactForm, FollowUpCompleteForm, FollowUpForm, LeadForm, LeadImportForm, LeadRemarkForm, LeadUpdateForm, OwnLeadForm, StudentApplicationForm, StudentDocumentForm, StudentSignupForm
from .models import BulkDataBatch, CRMApiKey, ChatMessage, College, CollegeCategory, Communication, Course, FollowUp, Lead, LeadActivity, LeadImportBatch, LeadRemark, NurtureLog, OnlineCourse, DistanceOnlineEducation, SocialClick, SocialLink, StudentApplication, UserProfile, UniversityMaster, StateMaster, CountryMaster
from .bulk_data import DATASETS, dataset_records, make_tabular_response, process_bulk_rows, read_image_zip, read_rows

def crm_role(user):
    if user.is_superuser: return "Admin"
    try: return user.crm_profile.role if user.crm_profile.active else ""
    except Exception: return "Staff" if user.is_staff else ""

def crm_allowed(user): return user.is_authenticated and bool(crm_role(user))

def bulk_data_allowed(user):
    return user.is_authenticated and (user.is_superuser or crm_role(user) in {"Admin", "Super Admin", "Admission Manager", "Staff"})

def scoped_leads(user):
    leads = Lead.objects.select_related("course", "preferred_college", "assigned_to", "uploaded_by")
    role = crm_role(user)
    if role in {"Admin", "Super Admin", "Admission Manager", "Staff"}: return leads
    if role == "College":
        try: return leads.filter(preferred_college=user.crm_profile.college)
        except Exception: return leads.none()
    if role == "Data Entry Operator": return leads.filter(uploaded_by=user)
    return leads.filter(Q(assigned_to=user) | Q(uploaded_by=user)).distinct()

def normalise_phone(value):
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else digits

def apply_lead_filters(queryset, params):
    """Apply the shared CRM report/export dimensions to a lead queryset."""
    filters = {
        "source": "source", "status": "status", "lead_category": "lead_category", "city": "city__iexact",
        "course": "course_id", "college": "preferred_college_id",
        "counsellor": "assigned_to_id",
    }
    for parameter, lookup in filters.items():
        value = params.get(parameter, "").strip()
        if value: queryset = queryset.filter(**{lookup: value})
    date_from, date_to = parse_date(params.get("date_from", "")), parse_date(params.get("date_to", ""))
    if date_from: queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to: queryset = queryset.filter(created_at__date__lte=date_to)
    return queryset

def _save_public_lead(form, source):
    """Create or update one CRM lead by mobile number from any public lead form."""
    data = form.cleaned_data
    phone = normalise_phone(data.get("phone"))
    lead = Lead.objects.filter(phone=phone).first()
    created = lead is None
    if created:
        lead = Lead(phone=phone)
    for field in ["name", "father_name", "email", "city", "course", "preferred_college", "preferred_distance_online"]:
        value = data.get(field)
        if value not in (None, ""):
            setattr(lead, field, value)
    lead.source = source
    lead.consent = True
    lead.save()
    LeadActivity.objects.create(lead=lead, activity_type="Created" if created else "Note", description=f"Enquiry captured from {source} lead form.")
    return lead

def home(request):
    q = request.GET.get("q", "").strip()
    colleges = College.objects.filter(active=True).prefetch_related("courses")
    if q: colleges = colleges.filter(Q(name__icontains=q) | Q(university__name__icontains=q) | Q(city__icontains=q) | Q(courses__name__icontains=q)).distinct()
    form = LeadForm(request.POST or None)
    if request.method == "POST":
        consent = request.POST.get("consent") == "on"
        if form.is_valid() and consent:
            _save_public_lead(form, "Website")
            messages.success(request, "Thanks! Your details are saved in Lead CRM. Helpline: 9429692142, 9911445580, 9911442142.")
            return redirect("home")
        messages.error(request, "Please enter valid details and accept the contact consent.")
    return render(request, "admissions/home.html", {"form": form, "apply_form": form, "colleges": colleges.prefetch_related("categories", "courses")[:12], "courses": Course.objects.prefetch_related("college_categories", "course_categories")[:12], "college_categories": CollegeCategory.objects.order_by("name"), "distance_online": DistanceOnlineEducation.objects.filter(active=True).order_by("-featured", "name")[:8], "q": q})

def health(request):
    try:
        with connection.cursor() as cursor: cursor.execute("SELECT 1"); cursor.fetchone()
        return JsonResponse({"status":"ok", "database":"connected", "engine":connection.vendor})
    except Exception:
        return JsonResponse({"status":"error", "database":"unavailable"}, status=503)

def about(request): return render(request, "admissions/about.html")

def contact(request):
    form = LeadForm(request.POST or None)
    if request.method == "POST":
        consent = request.POST.get("consent") == "on"
        if form.is_valid() and consent:
            _save_public_lead(form, "Contact Form")
            messages.success(request, "Thanks! Your details are saved in Lead CRM. Helpline: 9429692142, 9911445580, 9911442142.")
            return redirect("contact")
        messages.error(request, "Please enter valid details and accept the contact consent.")
    return render(request, "admissions/contact.html", {"form": form, "apply_form": form})

def online_courses(request):
    items = DistanceOnlineEducation.objects.filter(active=True).select_related("university", "state", "country").prefetch_related("categories")
    return render(request, "admissions/distance_online.html", {"items": items})

def courses_page(request):
    category = request.GET.get("category", "")
    q = request.GET.get("q", "").strip()
    courses = Course.objects.all().order_by("category", "name")
    if category: courses = courses.filter(category=category)
    if q: courses = courses.filter(Q(name__icontains=q) | Q(short_description__icontains=q) | Q(career_description__icontains=q) | Q(category__icontains=q))
    return render(request, "admissions/courses.html", {"course_list": courses, "categories": Course.CATEGORIES, "selected_category": category, "q": q})

def colleges_page(request):
    q = request.GET.get("q", "").strip()
    category = request.GET.get("category", "").strip()
    colleges = list(
        College.objects.filter(active=True)
        .select_related("university", "state", "country")
        .prefetch_related("categories")
        .order_by("name")
    )
    if category:
        colleges = [c for c in colleges if any(category.lower() in x.name.lower() for x in c.categories.all())]
    if q:
        college_ids = College.objects.filter(active=True).filter(
            Q(name__icontains=q) | Q(university__name__icontains=q) |
            Q(state__name__icontains=q) | Q(country__name__icontains=q) |
            Q(categories__name__icontains=q) | Q(short_description__icontains=q) |
            Q(description__icontains=q)
        ).values_list("pk", flat=True).distinct()
        colleges = [c for c in colleges if c.pk in set(college_ids)]

    # Courses are connected to colleges through their shared College Categories.
    # This keeps College CRM simple while Course CRM controls which college
    # categories each course belongs to.
    all_courses = Course.objects.prefetch_related("college_categories").order_by("name")
    category_course_map = {}
    for course in all_courses:
        for category_id in course.college_categories.values_list("id", flat=True):
            category_course_map.setdefault(category_id, []).append(course)
    for college in colleges:
        seen = set()
        connected = []
        for category in college.categories.all():
            for course in category_course_map.get(category.id, []):
                if course.pk not in seen:
                    seen.add(course.pk)
                    connected.append(course)
        college.connected_courses = connected

    return render(request, "admissions/colleges.html", {
        "college_list": colleges,
        "q": q,
        "selected_category": category,
    })

def college_detail(request, pk):
    college = get_object_or_404(College.objects.prefetch_related("categories", "state__country"), pk=pk, active=True)
    category_ids = college.categories.values_list("id", flat=True)
    connected_courses = Course.objects.filter(college_categories__id__in=category_ids).distinct().order_by("name")
    return render(request, "admissions/college_detail.html", {"college": college, "connected_courses": connected_courses})

def course_detail(request, pk):
    course = get_object_or_404(Course, pk=pk)
    related_colleges = College.objects.filter(active=True, courses=course).order_by("name")[:12]
    return render(request, "admissions/course_detail.html", {"course": course, "related_colleges": related_colleges})

def medical_colleges(request):
    # Medical colleges are now a category in the unified Colleges catalogue.
    return redirect("/colleges/?type=Medical")

def quick_apply(request):
    if request.method != "POST": return redirect("home")
    form = LeadForm(request.POST)
    consent = request.POST.get("consent") in {"on", "true", "1", "yes"}
    allowed_sources = {"Website", "AI Chatbot", "WhatsApp Chatbot", "Facebook Messenger", "Instagram Messenger", "Contact Form", "YouTube", "Facebook", "Instagram"}
    source = request.POST.get("lead_source", "Website")
    if source not in allowed_sources:
        source = "Website"
    if form.is_valid() and consent:
        _save_public_lead(form, source)
        messages.success(request, "Thanks! Your details are saved in Lead CRM. Helpline: 9429692142, 9911445580, 9911442142.")
    else:
        messages.error(request, "Please enter valid details and accept the contact consent.")
    return redirect(request.META.get("HTTP_REFERER") or "home")

def social_redirect(request, platform):
    link = get_object_or_404(SocialLink.objects.exclude(url=""), platform__iexact=platform, active=True)
    SocialClick.objects.create(platform=link.platform, source_page=request.META.get("HTTP_REFERER", "")[:200])
    return redirect(link.url)

def chatbot(request):
    if request.method != "POST": return JsonResponse({"error": "POST required"}, status=405)
    try: message = json.loads(request.body).get("message", "").strip()
    except (json.JSONDecodeError, AttributeError): message = ""
    if not message: return JsonResponse({"error": "Please enter a message."}, status=400)
    text = message.lower()
    if any(x in text for x in ["hello", "hi", "namaste"]): answer = "Hello! I can help with colleges, courses, eligibility, applications and documents. Which course are you interested in?"
    elif any(x in text for x in ["b.tech", "btech", "engineering"]): answer = "We provide B.Tech guidance for CSE, AI, ML, Data Science, Cyber Security, ECE, Mechanical and other branches. Please share your city and JEE rank through Free Counselling."
    elif any(x in text for x in ["bba", "bca", "bajmc", "law", "mba", "pharma", "medical"]): answer = "We can compare colleges, eligibility, fees and application options for this course. Use Free Counselling or WhatsApp to receive a personalised shortlist."
    elif any(x in text for x in ["document", "marksheet", "aadhaar", "certificate"]): answer = "Students can upload photograph, signature, Aadhaar, Class 10 and 12 marksheets, entrance scorecard and certificates from My Application."
    elif any(x in text for x in ["fee", "scholarship", "cost"]): answer = "Fees and scholarships depend on the college and course. Scholarship assistance is available. Share your course and marks for suitable options."
    elif any(x in text for x in ["online", "distance"]): answer = "Open the Online Courses page to explore available programmes. You can submit an enquiry for counsellor guidance."
    elif any(x in text for x in ["contact", "phone", "whatsapp", "counsellor"]): answer = "Please use the WhatsApp button or Free Counselling form. A counsellor will contact you during working hours."
    else: answer = "I can help with course selection, college comparison, eligibility, fees, scholarships, applications and documents. Please tell me your preferred course and city."
    if not request.session.session_key: request.session.create()
    ChatMessage.objects.create(session_key=request.session.session_key, user_message=message, bot_response=answer)
    return JsonResponse({"reply": answer})

def student_signup(request):
    form = StudentSignupForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save(); login(request, user); messages.success(request, "Account created. Complete your application below."); return redirect("student_dashboard")
    return render(request, "registration/signup.html", {"form": form})

def crm_login(request):
    form=AuthenticationForm(request,data=request.POST or None)
    if request.method=="POST" and form.is_valid():
        login(request,form.get_user())
        request.session.set_expiry(1209600 if request.POST.get("remember_me") else 0)
        return redirect("dashboard" if crm_allowed(form.get_user()) else "student_dashboard")
    return render(request,"registration/login.html",{"form":form})

@login_required
def student_dashboard(request):
    if crm_allowed(request.user): return redirect("dashboard")
    application, _ = StudentApplication.objects.get_or_create(user=request.user, defaults={"full_name": request.user.get_full_name() or request.user.username, "phone": "", "email": request.user.email})
    app_form = StudentApplicationForm(request.POST or None, instance=application, prefix="application")
    doc_form = StudentDocumentForm(request.POST or None, request.FILES or None, prefix="document")
    if request.method == "POST" and "save_application" in request.POST and app_form.is_valid():
        app_form.save(); messages.success(request, "Application details saved."); return redirect("student_dashboard")
    if request.method == "POST" and "submit_application" in request.POST and app_form.is_valid():
        item = app_form.save(commit=False); item.status = "Submitted"; item.submitted_at = timezone.now(); item.save(); messages.success(request, "Application submitted successfully."); return redirect("student_dashboard")
    if request.method == "POST" and "upload_document" in request.POST and doc_form.is_valid():
        document = doc_form.save(commit=False); document.application = application; document.save(); messages.success(request, "Document uploaded."); return redirect("student_dashboard")
    return render(request, "admissions/student_dashboard.html", {"application": application, "app_form": app_form, "doc_form": doc_form})

@login_required
def add_own_lead(request):
    if not crm_allowed(request.user):
        return redirect("student_dashboard")
    profile = getattr(request.user, "crm_profile", None)
    if not (request.user.is_superuser or (profile and profile.can_add_own_leads)):
        return HttpResponse("Permission to add own leads is required", status=403)
    form = OwnLeadForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        phone = normalise_phone(form.cleaned_data["phone"])
        existing = Lead.objects.filter(phone=phone).first()
        if existing:
            messages.warning(request, "This mobile number already exists in Lead CRM.")
            return redirect("lead_detail", pk=existing.pk) if scoped_leads(request.user).filter(pk=existing.pk).exists() else redirect("dashboard")
        lead = form.save(commit=False)
        lead.phone = phone
        lead.uploaded_by = request.user
        lead.assigned_to = request.user
        lead.consent = True
        lead.save()
        LeadActivity.objects.create(lead=lead, activity_type="Created", description="Lead added by user to own CRM panel.", created_by=request.user)
        messages.success(request, "Lead added to your panel.")
        return redirect("lead_detail", pk=lead.pk)
    return render(request, "admissions/add_own_lead.html", {"form": form})

@login_required
def dashboard(request):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    leads = scoped_leads(request.user).order_by("-created_at")
    status = request.GET.get("status", "")
    if status: leads = leads.filter(status=status)
    stats = dict(scoped_leads(request.user).values_list("status").annotate(total=Count("id")))
    due = request.user.followup_set.filter(completed=False, due_at__lte=timezone.now()).select_related("lead") if hasattr(request.user, "followup_set") else []
    featured = College.objects.filter(active=True, featured=True).exclude(image="")[:4]
    profile = getattr(request.user, "crm_profile", None)
    can_import = request.user.is_superuser or crm_role(request.user) in {"Admin", "Super Admin", "Admission Manager", "Staff", "Data Entry Operator", "Consultant"} or bool(profile and profile.can_import_leads)
    can_export = request.user.is_superuser or crm_role(request.user) in {"Admin", "Super Admin", "Admission Manager", "Staff"} or bool(profile and profile.can_export_leads)
    base = scoped_leads(request.user)
    today = timezone.localdate()
    today_leads = base.filter(created_at__date=today).count()
    followups_due = FollowUp.objects.filter(lead__in=base, completed=False, due_at__lte=timezone.now()).count()
    visits_today = base.filter(status__in=["Visit Planned", "Counselling Booked"]).count()
    applications = base.filter(status__in=["Application Started", "Application Submitted", "Documents Pending"]).count()
    admissions = base.filter(status="Admission Confirmed").count()
    unassigned = base.filter(assigned_to__isnull=True).count()
    source_stats = list(base.values("source").annotate(total=Count("id")).order_by("-total")[:6])
    counsellor_stats = list(base.filter(assigned_to__isnull=False).values("assigned_to__username").annotate(total=Count("id"), admissions=Count("id", filter=Q(status="Admission Confirmed"))).order_by("-admissions", "-total")[:8])
    attention = {"overdue": followups_due, "unassigned": unassigned, "documents": base.filter(status="Documents Pending").count(), "fees": base.filter(status="Fee Pending").count(), "applications": base.filter(status="Application Started").count()}
    return render(request, "admissions/dashboard.html", {"leads": leads[:100], "stats": stats, "statuses": Lead.STATUS, "selected": status, "due": due, "featured": featured, "crm_role": crm_role(request.user), "can_import": can_import, "can_export": can_export, "can_bulk_data": bulk_data_allowed(request.user), "is_crm_admin": admin_authorised(request.user), "today_leads":today_leads, "followups_due":followups_due, "visits_today":visits_today, "applications_count":applications, "admissions_count":admissions, "attention":attention, "source_stats":source_stats, "counsellor_stats":counsellor_stats})

@login_required
def bulk_data_manager(request):
    if not bulk_data_allowed(request.user): return HttpResponse("Bulk data permission required", status=403)
    form = BulkDataImportForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        upload = form.cleaned_data["file"]
        try:
            rows = read_rows(upload)
            if not rows: raise ValueError("The uploaded file has no data rows.")
            images = read_image_zip(form.cleaned_data.get("image_zip"))
            batch = process_bulk_rows(form.cleaned_data["dataset"], rows, upload.name, request.user, form.cleaned_data["mode"], images)
            messages.success(request, f"Import complete: {batch.created_count} created, {batch.updated_count} updated, {batch.skipped_count} skipped, {batch.rejected_count} rejected and {batch.imported_image_count} images saved.")
            return redirect("bulk_data_manager")
        except Exception as exc:
            messages.error(request, f"Import failed: {exc}")
    batches = BulkDataBatch.objects.select_related("uploaded_by").order_by("-created_at")[:30]
    return render(request, "admissions/bulk_data.html", {"form": form, "datasets": DATASETS, "batches": batches})

@login_required
def bulk_data_template(request, dataset, file_format):
    if not bulk_data_allowed(request.user): return HttpResponse("Bulk data permission required", status=403)
    if dataset not in DATASETS or file_format not in {"csv", "xlsx"}: return HttpResponse("Invalid template", status=404)
    return make_tabular_response(dataset, file_format, template=True)

@login_required
def bulk_data_export(request, dataset, file_format):
    if not bulk_data_allowed(request.user): return HttpResponse("Bulk data permission required", status=403)
    if dataset not in DATASETS or file_format not in {"csv", "xlsx"}: return HttpResponse("Invalid export", status=404)
    records = scoped_leads(request.user).order_by("-created_at") if dataset == "leads" else dataset_records(dataset)
    return make_tabular_response(dataset, file_format, records=records)

@login_required
def bulk_data_errors(request, pk):
    if not bulk_data_allowed(request.user): return HttpResponse("Bulk data permission required", status=403)
    batch = get_object_or_404(BulkDataBatch, pk=pk)
    response = HttpResponse(content_type="text/csv; charset=utf-8"); response.write("\ufeff")
    response["Content-Disposition"] = f'attachment; filename="bulk-errors-{batch.pk}.csv"'
    writer = csv.writer(response); writer.writerow(["Row", "Matching Key", "Error"])
    try: errors = json.loads(batch.error_report or "[]")
    except json.JSONDecodeError: errors = [{"row":"", "key":"", "error":batch.error_report}]
    for item in errors: writer.writerow([item.get("row", ""), item.get("key", ""), item.get("error", "")])
    return response

@login_required
def lead_detail(request, pk):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    lead = get_object_or_404(scoped_leads(request.user), pk=pk)
    update_form, follow_form, communication_form, remark_form = LeadUpdateForm(instance=lead), FollowUpForm(), CommunicationForm(), LeadRemarkForm()
    if crm_role(request.user) not in {"Admin", "Super Admin", "Admission Manager", "Staff"}: update_form.fields.pop("assigned_to", None)
    if request.method == "POST":
        if "update_lead" in request.POST:
            old_stage, old_assignee = lead.status, lead.assigned_to
            update_form = LeadUpdateForm(request.POST, instance=lead)
            if crm_role(request.user) not in {"Admin", "Super Admin", "Admission Manager", "Staff"}: update_form.fields.pop("assigned_to", None)
            if update_form.is_valid():
                updated = update_form.save()
                if updated.call_outcome != "Untouched": updated.last_contacted_at = timezone.now(); updated.save(update_fields=["last_contacted_at", "updated_at"])
                if old_stage != updated.status: LeadActivity.objects.create(lead=lead, activity_type="Stage Changed", description=f"Pipeline moved from {old_stage} to {updated.status}", old_stage=old_stage, new_stage=updated.status, created_by=request.user)
                if old_assignee != updated.assigned_to: LeadActivity.objects.create(lead=lead, activity_type="Assignment", description=f"Lead assigned to {updated.assigned_to or 'Unassigned'}", created_by=request.user)
                LeadActivity.objects.create(lead=lead, activity_type="Call", description=f"Outcome: {updated.call_outcome}. Next action: {updated.next_action or 'Not set'}", created_by=request.user)
                messages.success(request, "Lead progress updated."); return redirect("lead_detail", pk=pk)
        elif "add_followup" in request.POST:
            follow_form = FollowUpForm(request.POST)
            if follow_form.is_valid():
                item = follow_form.save(commit=False); item.lead, item.created_by = lead, request.user; item.save()
                LeadActivity.objects.create(lead=lead, activity_type="Follow-up", description=f"Scheduled for {item.due_at}: {item.note}", created_by=request.user)
                messages.success(request, "Follow-up scheduled."); return redirect("lead_detail", pk=pk)
        elif "save_communication" in request.POST or "send_communication" in request.POST:
            communication_form = CommunicationForm(request.POST)
            if communication_form.is_valid():
                item = communication_form.save(commit=False); item.lead, item.created_by = lead, request.user
                item.status = "Queued" if "send_communication" in request.POST else "Draft"; item.save()
                if "send_communication" in request.POST: deliver_communication(item)
                LeadActivity.objects.create(lead=lead, activity_type="Communication", description=f"{item.channel} {item.status.lower()}: {item.subject or item.message[:80]}", created_by=request.user)
                messages.success(request, f"{item.channel} {item.status.lower()} saved."); return redirect("lead_detail", pk=pk)
        elif "add_remark" in request.POST:
            remark_form = LeadRemarkForm(request.POST)
            if remark_form.is_valid():
                item = remark_form.save(commit=False); item.lead, item.created_by = lead, request.user; item.save()
                LeadActivity.objects.create(lead=lead, activity_type="Note", description=item.remark, created_by=request.user)
                messages.success(request, "Remark added."); return redirect("lead_detail", pk=pk)
        else:
            follow_form = FollowUpForm(request.POST)
        if follow_form.is_valid() and "add_followup" in request.POST:
            item = follow_form.save(commit=False); item.lead, item.created_by = lead, request.user; item.save()
            messages.success(request, "Follow-up scheduled."); return redirect("lead_detail", pk=pk)
    application = StudentApplication.objects.filter(Q(phone=lead.phone) | (Q(email__iexact=lead.email) if lead.email else Q(pk__isnull=True))).select_related("course", "college", "user").prefetch_related("documents", "payments").first()
    return render(request, "admissions/lead_detail.html", {"lead": lead, "update_form": update_form, "follow_form": follow_form, "communication_form": communication_form, "remark_form":remark_form, "pipeline_stages":Lead.STATUS, "application":application})

@login_required
def complete_followup(request, pk):
    item = get_object_or_404(FollowUp.objects.filter(lead__in=scoped_leads(request.user)), pk=pk)
    if request.method == "POST":
        form = FollowUpCompleteForm(request.POST)
        if form.is_valid():
            item.outcome, item.completed, item.completed_at = form.cleaned_data["outcome"], True, timezone.now(); item.save(update_fields=["outcome", "completed", "completed_at"])
            LeadActivity.objects.create(lead=item.lead, activity_type="Follow-up", description=f"Completed: {item.outcome}", created_by=request.user)
            messages.success(request, "Follow-up completed and added to lead history.")
    return redirect("lead_detail", pk=item.lead_id)

def deliver_communication(item):
    try:
        if item.channel == "Email":
            if not item.lead.email: raise ValueError("Lead has no email address")
            send_mail(item.subject or "College Admission Update", item.message, settings.DEFAULT_FROM_EMAIL, [item.lead.email], fail_silently=False)
            response = "Accepted by configured email backend"
        elif item.channel == "WhatsApp":
            if not settings.WHATSAPP_API_URL or not settings.WHATSAPP_ACCESS_TOKEN: raise ValueError("WhatsApp Business API is not configured")
            payload = {"messaging_product":"whatsapp", "to":f"91{normalise_phone(item.lead.phone)}", "type":"text", "text":{"body":item.message}}
            req = urllib.request.Request(settings.WHATSAPP_API_URL, data=json.dumps(payload).encode(), headers={"Authorization":f"Bearer {settings.WHATSAPP_ACCESS_TOKEN}", "Content-Type":"application/json"}, method="POST")
            response = urllib.request.urlopen(req, timeout=15).read().decode()
        else:
            if not settings.SMS_API_URL or not settings.SMS_AUTH_KEY: raise ValueError("SMS API is not configured")
            payload = {"mobile":f"91{normalise_phone(item.lead.phone)}", "message":item.message, "sender":settings.SMS_SENDER_ID}
            req = urllib.request.Request(settings.SMS_API_URL, data=json.dumps(payload).encode(), headers={"authkey":settings.SMS_AUTH_KEY, "Content-Type":"application/json"}, method="POST")
            response = urllib.request.urlopen(req, timeout=15).read().decode()
        item.status, item.sent_at, item.provider_response = "Sent", timezone.now(), response[:2000]
    except Exception as exc:
        item.status, item.provider_response = "Failed", str(exc)[:2000]
    item.save(update_fields=["status", "sent_at", "provider_response"])

@login_required
def import_leads(request):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    profile = getattr(request.user, "crm_profile", None)
    allowed = request.user.is_superuser or crm_role(request.user) in {"Admin", "Super Admin", "Admission Manager", "Staff", "Data Entry Operator", "Consultant"} or bool(profile and profile.can_import_leads)
    if not allowed: return HttpResponse("Import permission required", status=403)
    form = LeadImportForm(request.POST or None, request.FILES or None)
    form.fields["assigned_to"].queryset = form.fields["assigned_to"].queryset.filter(Q(is_staff=True) | Q(crm_profile__role__in=["Employee", "Agent", "Consultant"])).distinct()
    if request.method == "POST" and form.is_valid():
        upload = form.cleaned_data["file"]
        try:
            rows = read_lead_rows(upload)
            batch = process_lead_rows(rows, upload.name, request.user, form.cleaned_data["assigned_to"], form.cleaned_data["source"])
            messages.success(request, f"Import complete: {batch.created_count} created, {batch.updated_count} updated, {batch.rejected_count} rejected.")
            return redirect("dashboard")
        except Exception as exc: messages.error(request, f"Import failed: {exc}")
    return render(request, "admissions/import_leads.html", {"form": form, "recent_batches": LeadImportBatch.objects.filter(uploaded_by=request.user).order_by("-created_at")[:10]})

def read_lead_rows(upload):
    if upload.name.lower().endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    book = load_workbook(upload, read_only=True, data_only=True); sheet = book.active
    values = list(sheet.iter_rows(values_only=True))
    if not values: return []
    headers = [str(x or "").strip() for x in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]]

def process_lead_rows(rows, file_name, user, assigned_to, source):
    aliases = {"name":["Student Name","Name","Student"], "father_name":["Father Name","Father"], "phone":["Mobile Number","Mobile","Phone","Phone Number"], "email":["Email","Email Address"], "city":["City","Location"], "course":["Course","Interested Course"], "score":["JEE / CLAT","CET/JEE/NEET Rank","Score","Rank"]}
    def pick(row, key):
        lowered = {str(k).strip().lower(): v for k,v in row.items()}
        return next((lowered.get(x.lower()) for x in aliases[key] if lowered.get(x.lower()) not in (None, "")), "")
    batch = LeadImportBatch.objects.create(uploaded_by=user, assigned_to=assigned_to, file_name=file_name, total_rows=len(rows)); errors=[]
    for number, row in enumerate(rows, start=2):
        phone, name = normalise_phone(pick(row,"phone")), str(pick(row,"name") or "").strip()
        if len(phone) != 10 or not name: batch.rejected_count += 1; errors.append(f"Row {number}: valid name and 10-digit mobile required"); continue
        course_name = str(pick(row,"course") or "").strip(); course = Course.objects.filter(name__iexact=course_name).first() if course_name else None
        values = {"name":name, "father_name":str(pick(row,"father_name") or "").strip(), "email":str(pick(row,"email") or "").strip(), "city":str(pick(row,"city") or "").strip(), "score":str(pick(row,"score") or "").strip(), "source":source, "course":course}
        lead = Lead.objects.filter(phone=phone).first()
        if lead:
            for field, value in values.items():
                if value not in (None, ""): setattr(lead, field, value)
            if assigned_to and crm_role(user) in {"Admin", "Super Admin", "Admission Manager", "Staff"}: lead.assigned_to = assigned_to
            lead.save(); batch.updated_count += 1
        else:
            Lead.objects.create(phone=phone, assigned_to=assigned_to, uploaded_by=user, consent=True, **values)
            batch.created_count += 1
    batch.error_report = "\n".join(errors); batch.save(); return batch

@login_required
def export_leads(request, file_format):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    profile = getattr(request.user, "crm_profile", None)
    allowed = request.user.is_superuser or crm_role(request.user) in {"Admin", "Super Admin", "Admission Manager", "Staff"} or bool(profile and profile.can_export_leads)
    if not allowed: return HttpResponse("Export permission required", status=403)
    rows = apply_lead_filters(scoped_leads(request.user), request.GET).order_by("-created_at")
    if request.GET.get("preset") == "today": rows = rows.filter(created_at__date=timezone.localdate())
    if request.GET.get("preset") == "hot": rows = rows.filter(status="Hot")
    if request.GET.get("preset") == "admission": rows = rows.filter(status__in=["Admission Confirmed", "Registered"])
    if request.GET.get("counsellor"): rows = rows.filter(assigned_to_id=request.GET["counsellor"])
    if request.GET.get("course"): rows = rows.filter(course_id=request.GET["course"])
    if request.GET.get("college"): rows = rows.filter(preferred_college_id=request.GET["college"])
    headers = ["Student Name","Mobile Number","Email","City","Course","Preferred College","Source","Status","Assigned To","Created"]
    if file_format == "pdf":
        response=HttpResponse(content_type="application/pdf"); response["Content-Disposition"]='attachment; filename="crm-leads.pdf"'
        pdf=canvas.Canvas(response,pagesize=landscape(A4)); width,height=landscape(A4); y=height-35
        pdf.setFont("Helvetica-Bold",14); pdf.drawString(30,y,"College Admission CRM - Lead Report"); y-=24; pdf.setFont("Helvetica",8)
        for item in rows[:1000]:
            pdf.drawString(30,y,f"{item.pk} | {item.name[:24]} | {item.phone} | {str(item.course or '-')[:18]} | {item.status[:20]} | {str(item.assigned_to or '-')[:16]}"); y-=13
            if y<30: pdf.showPage(); pdf.setFont("Helvetica",8); y=height-30
        pdf.save(); return response
    if file_format == "csv":
        response = HttpResponse(content_type="text/csv"); response["Content-Disposition"] = 'attachment; filename="crm-leads.csv"'
        writer = csv.writer(response); writer.writerow(headers)
        for x in rows: writer.writerow([x.name,x.phone,x.email,x.city,x.course or "",x.preferred_college or "",x.source,x.status,x.assigned_to or "",x.created_at.strftime("%Y-%m-%d %H:%M")])
        return response
    book=Workbook(); sheet=book.active; sheet.title="CRM Leads"; sheet.append(headers)
    for x in rows: sheet.append([x.name,x.phone,x.email,x.city,str(x.course or ""),str(x.preferred_college or ""),x.source,x.status,str(x.assigned_to or ""),x.created_at.strftime("%Y-%m-%d %H:%M")])
    stream=io.BytesIO(); book.save(stream); response=HttpResponse(stream.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); response["Content-Disposition"]='attachment; filename="crm-leads.xlsx"'; return response

@login_required
def reports_dashboard(request):
    if crm_role(request.user) not in {"Admin", "Super Admin", "Admission Manager", "Staff"}: return HttpResponse("Report permission required", status=403)
    now, today = timezone.now(), timezone.localdate()
    admission_stages = ["Admission Confirmed", "Registered"]
    leads = apply_lead_filters(scoped_leads(request.user), request.GET)
    total = leads.count()
    admissions = leads.filter(status__in=admission_stages).count()
    contacted = leads.exclude(call_outcome="Untouched").count()
    followups = FollowUp.objects.filter(lead__in=leads)

    source_data = list(leads.values("source").annotate(
        total=Count("id", distinct=True),
        admissions=Count("id", filter=Q(status__in=admission_stages), distinct=True),
    ).order_by("-total"))
    for row in source_data:
        row["conversion"] = round(row["admissions"] * 100 / row["total"], 1) if row["total"] else 0

    counsellors = list(leads.exclude(assigned_to=None).values("assigned_to__username").annotate(
        total=Count("id", distinct=True),
        admissions=Count("id", filter=Q(status__in=admission_stages), distinct=True),
        pending=Count("followups", filter=Q(followups__completed=False), distinct=True),
        completed=Count("followups", filter=Q(followups__completed=True), distinct=True),
    ).order_by("-admissions", "-total"))
    for row in counsellors:
        row["conversion"] = round(row["admissions"] * 100 / row["total"], 1) if row["total"] else 0

    nurture_logs = NurtureLog.objects.filter(lead__in=leads)
    communications = Communication.objects.filter(lead__in=leads)
    nurture_data = nurture_logs.values("day", "status").annotate(total=Count("id")).order_by("day", "status")
    communication_data = communications.values("channel", "status").annotate(total=Count("id")).order_by("channel", "status")
    pending_followups = followups.filter(completed=False)
    context = {
        "total": total, "today": leads.filter(created_at__date=today).count(),
        "hot": leads.filter(status="Hot").count(), "admissions": admissions,
        "conversion_rate": round(admissions * 100 / total, 1) if total else 0,
        "contact_rate": round(contacted * 100 / total, 1) if total else 0,
        "total_campaigns": nurture_logs.count(), "total_communications": communications.count(),
        "pending": pending_followups.count(), "completed_followups": followups.filter(completed=True).count(),
        "overdue": pending_followups.filter(due_at__lt=now).count(),
        "due_today": pending_followups.filter(due_at__date=today).count(),
        "upcoming": pending_followups.filter(due_at__gt=now).exclude(due_at__date=today).count(),
        "pipeline_data": leads.values("status").annotate(total=Count("id")).order_by("-total"),
        "outcome_data": leads.values("call_outcome").annotate(total=Count("id")).order_by("-total"),
        "source_data": source_data,
        "course_data": leads.values("course__name").annotate(total=Count("id"), admissions=Count("id", filter=Q(status__in=admission_stages))).order_by("-total")[:15],
        "college_data": leads.values("preferred_college__name").annotate(total=Count("id"), admissions=Count("id", filter=Q(status__in=admission_stages))).order_by("-total")[:15],
        "city_data": leads.values("city").annotate(total=Count("id")).order_by("-total")[:15],
        "daily_data": leads.annotate(day=TruncDate("created_at")).values("day").annotate(total=Count("id"), admissions=Count("id", filter=Q(status__in=admission_stages))).order_by("-day")[:31],
        "nurture_data": nurture_data, "communication_data": communication_data,
        "counsellors": counsellors,
        "sources": Lead.SOURCE, "statuses": Lead.STATUS,
        "courses": Course.objects.order_by("name"), "colleges": College.objects.order_by("name"),
        "users": User.objects.filter(Q(is_staff=True) | Q(crm_profile__active=True)).distinct().order_by("username"),
        "cities": scoped_leads(request.user).exclude(city="").values_list("city", flat=True).distinct().order_by("city"),
        "querystring": request.GET.urlencode(),
    }
    return render(request,"admissions/reports.html",context)

@login_required
def followup_calendar(request):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    items = FollowUp.objects.filter(lead__in=scoped_leads(request.user)).select_related("lead", "lead__assigned_to").order_by("due_at")
    month = request.GET.get("month", "")
    if month:
        try:
            year, number = [int(x) for x in month.split("-", 1)]; items = items.filter(due_at__year=year, due_at__month=number)
        except (TypeError, ValueError): pass
    return render(request, "admissions/followup_calendar.html", {"followups":items[:500], "selected_month":month})

def admin_authorised(user): return user.is_authenticated and (user.is_superuser or crm_role(user) in {"Admin", "Super Admin"})

@login_required
def consultant_panel(request):
    if not admin_authorised(request.user): return HttpResponse("Admin authorisation required", status=403)
    form = ConsultantCreateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save(); messages.success(request, "Consultant account authorised and created."); return redirect("consultant_panel")
    consultants = UserProfile.objects.filter(role="Consultant").select_related("user").order_by("-user__date_joined")
    return render(request, "admissions/consultant_panel.html", {"form":form, "consultants":consultants})

@login_required
def api_key_panel(request):
    if not admin_authorised(request.user): return HttpResponse("Admin authorisation required", status=403)
    form, generated_key = ApiKeyCreateForm(request.POST or None), None
    form.fields["owner"].queryset = User.objects.filter(Q(is_superuser=True) | Q(crm_profile__active=True)).distinct()
    if request.method == "POST" and form.is_valid():
        generated_key = "ca_live_" + secrets.token_urlsafe(32)
        CRMApiKey.objects.create(owner=form.cleaned_data["owner"], name=form.cleaned_data["name"], prefix=generated_key[:16], key_hash=hashlib.sha256(generated_key.encode()).hexdigest(), expires_at=form.cleaned_data["expires_at"], created_by=request.user)
        form = ApiKeyCreateForm(); form.fields["owner"].queryset = User.objects.filter(Q(is_superuser=True) | Q(crm_profile__active=True)).distinct()
    keys = CRMApiKey.objects.select_related("owner").order_by("-created_at")
    return render(request, "admissions/api_keys.html", {"form":form, "keys":keys, "generated_key":generated_key})

@login_required
def revoke_api_key(request, pk):
    if not admin_authorised(request.user): return HttpResponse("Admin authorisation required", status=403)
    if request.method == "POST": CRMApiKey.objects.filter(pk=pk).update(active=False); messages.success(request, "API key revoked.")
    return redirect("api_key_panel")

def authenticate_api_key(request):
    raw = request.headers.get("X-API-Key", "")
    if not raw.startswith("ca_live_"): return None
    digest, now = hashlib.sha256(raw.encode()).hexdigest(), timezone.now()
    key = CRMApiKey.objects.filter(prefix=raw[:16], active=True).select_related("owner").first()
    if not key or not hmac.compare_digest(key.key_hash, digest) or (key.expires_at and key.expires_at <= now): return None
    key.last_used_at = now; key.save(update_fields=["last_used_at"]); return key

@csrf_exempt
def api_leads(request):
    if request.method != "POST": return JsonResponse({"error":"POST required"}, status=405)
    api_key = authenticate_api_key(request)
    if not api_key: return JsonResponse({"error":"Invalid, expired or revoked API key"}, status=401)
    try: payload = json.loads(request.body)
    except (json.JSONDecodeError, UnicodeDecodeError): return JsonResponse({"error":"Valid JSON required"}, status=400)
    phone, name = normalise_phone(payload.get("mobile") or payload.get("phone")), str(payload.get("name") or "").strip()
    if len(phone) != 10 or not name: return JsonResponse({"error":"name and valid 10-digit mobile are required"}, status=400)
    course = Course.objects.filter(name__iexact=str(payload.get("course") or "").strip()).first()
    lead = Lead.objects.filter(phone=phone).first(); created = lead is None
    if created: lead = Lead(phone=phone, uploaded_by=api_key.owner, assigned_to=api_key.owner, consent=bool(payload.get("consent", False)))
    for field, value in {"name":name, "email":payload.get("email", ""), "city":payload.get("city", ""), "score":payload.get("score", ""), "source":payload.get("source", "Website")}.items():
        if value not in (None, ""): setattr(lead, field, value)
    if course: lead.course = course
    lead.save()
    LeadActivity.objects.create(lead=lead, activity_type="Created" if created else "Note", description=f"Lead {'created' if created else 'updated'} through CRM API key {api_key.prefix}…", created_by=api_key.owner)
    return JsonResponse({"success":True, "created":created, "lead_id":lead.pk, "mobile":lead.phone}, status=201 if created else 200)

@login_required
def integration_settings(request):
    if not admin_authorised(request.user): return HttpResponse("Admin authorisation required", status=403)
    smtp_ok = all([settings.EMAIL_HOST, settings.EMAIL_HOST_USER, settings.EMAIL_HOST_PASSWORD]) and "console" not in settings.EMAIL_BACKEND
    whatsapp_ok = all([settings.WHATSAPP_API_URL, settings.WHATSAPP_ACCESS_TOKEN])
    sms_ok = all([settings.SMS_API_URL, settings.SMS_AUTH_KEY, settings.SMS_SENDER_ID])
    if request.method == "POST" and request.POST.get("test_smtp"):
        if not request.user.email: messages.error(request, "Add an email address to your Admin user first.")
        elif not smtp_ok: messages.error(request, "SMTP credentials are incomplete.")
        else:
            try:
                send_mail("College Admission CRM SMTP Test", "SMTP configuration is working.", settings.DEFAULT_FROM_EMAIL, [request.user.email], fail_silently=False)
                messages.success(request, f"Test email sent to {request.user.email}.")
            except Exception as exc: messages.error(request, f"SMTP test failed: {exc}")
        return redirect("integration_settings")
    return render(request, "admissions/integration_settings.html", {"smtp_ok":smtp_ok, "whatsapp_ok":whatsapp_ok, "sms_ok":sms_ok, "email_host":settings.EMAIL_HOST, "email_user":settings.EMAIL_HOST_USER, "whatsapp_url":settings.WHATSAPP_API_URL, "sms_url":settings.SMS_API_URL, "sms_sender":settings.SMS_SENDER_ID})

@login_required
def admission_pipeline(request):
    if not crm_allowed(request.user): return redirect("student_dashboard")
    leads = scoped_leads(request.user).order_by("-updated_at")
    stages = ["New Enquiry", "Contacted", "Interested", "Hot", "Counselling Booked", "Counselling Done", "Application Started", "Documents Pending", "Application Submitted", "Visit Planned", "Offer Issued", "Fee Pending", "Admission Confirmed"]
    columns = [{"name": stage, "items": leads.filter(status=stage)[:50], "count": leads.filter(status=stage).count()} for stage in stages]
    return render(request, "admissions/pipeline.html", {"columns":columns})

CRM_MODULES = [
 ('college','College CRM','College catalogue, university mapping, fees, approvals, ranking and media.'),('course','Course CRM','Courses, specialisations, eligibility, fees, seats and brochures.'),('online_course','Online Course CRM','Online programmes, LMS, accreditation, fees and applications.'),('student','Student CRM','Student applications, documents, payments and admission progress.'),('medical_college','Medical College CRM','Medical institutions, programmes, approvals and admission information.'),('social_seo','Social Media & SEO','SEO content, keywords, campaigns and social publishing workflow.'),('entrance_exam','Entrance Exam CRM','Entrance exams, important dates, eligibility and official links.'),('classroom','Classroom CRM','Classes, trainers, schedules, rooms and meeting links.'),('employee_hr','Employee & HR','Staff profiles, departments, managers and account status.'),('counsellor_calling','Counsellor & Calling','Counsellor queues, calls, outcomes and follow-up workflow.'),('leads','Lead & Enquiry CRM','Lead pipeline, assignments, stages, visits and conversions.'),('communications','Communication Hub','Email, WhatsApp, SMS and communication templates.'),('finance','Finance & Payments','Student fees, payment status, dues and refund tracking.'),('tasks','Tasks & Follow-ups','Team tasks, priorities, deadlines and follow-ups.'),('reports','Reports & Analytics','Admissions, lead source, conversion and counsellor reporting.'),('partners','Partner / Consultant','College partners, consultants and controlled CRM access.'),('roles','Roles & Permissions','Per-user module access and action-level permissions.')]

def module_access(user, key):
    if user.is_superuser or crm_role(user) in {'Admin','Super Admin'}: return True
    return user.module_permissions.filter(module=key, can_view=True).exists()

@login_required
def module_dashboard(request, module):
    item=next((x for x in CRM_MODULES if x[0]==module),None)
    if not item: return HttpResponse('Module not found',status=404)
    if not module_access(request.user,module): return HttpResponse('You do not have access to this CRM module.',status=403)
    counts={'Colleges':College.objects.count(),'Courses':Course.objects.count(),'Students':StudentApplication.objects.count(),'Leads':scoped_leads(request.user).count()}
    return render(request,'admissions/module_dashboard.html',{'module_key':module,'module_title':item[1],'module_description':item[2],'counts':counts,'is_crm_admin':admin_authorised(request.user)})
